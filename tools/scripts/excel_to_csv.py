import csv
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Chưa cài openpyxl.")
    print("Chạy: pip install openpyxl")
    raise SystemExit(1)

def main():
    if len(sys.argv) < 3:
        print('Cách dùng:')
        print('python excel_to_csv.py "input.xlsx" "output.csv"')
        raise SystemExit(1)

    input_file = Path(sys.argv[1])
    output_file = Path(sys.argv[2])

    wb = load_workbook(input_file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    with output_file.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])

    print(f"Đã chuyển: {input_file}")
    print(f"Thành: {output_file}")

if __name__ == "__main__":
    main()
